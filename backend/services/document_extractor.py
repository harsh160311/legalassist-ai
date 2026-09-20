"""
Document text extraction service.
Supports PDF, DOCX, and TXT files.
"""

import os
from pathlib import Path
from typing import Optional
import PyPDF2
from docx import Document


class DocumentExtractor:
    """Extracts text from various document formats."""

    @staticmethod
    def extract_text(file_path: str) -> dict:
        """
        Extract text from a document file.
        
        Args:
            file_path: Path to the document file
            
        Returns:
            dict with keys: text, pages (list of page texts), metadata
        """
        path = Path(file_path)
        
        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")
        
        extension = path.suffix.lower()
        
        if extension == ".pdf":
            return DocumentExtractor._extract_pdf(path)
        elif extension in [".docx", ".doc"]:
            return DocumentExtractor._extract_docx(path)
        elif extension == ".txt":
            return DocumentExtractor._extract_txt(path)
        else:
            raise ValueError(f"Unsupported file format: {extension}")

    @staticmethod
    def _extract_pdf(path: Path) -> dict:
        """Extract text from PDF file, including AcroForm field values."""
        pages = []
        metadata = {}
        form_fields = {}

        try:
            with open(path, "rb") as file:
                reader = PyPDF2.PdfReader(file)

                # Extract metadata
                if reader.metadata:
                    metadata = {
                        "title": reader.metadata.title or "",
                        "author": reader.metadata.author or "",
                        "pages": len(reader.pages)
                    }
                else:
                    metadata = {"pages": len(reader.pages)}

                # Extract AcroForm field values (filled-in form data)
                try:
                    fields = reader.get_fields()
                    if fields:
                        for field_name, field_data in fields.items():
                            value = field_data.get("/V")
                            if value and str(value).strip():
                                # Clean up the field name for readability
                                clean_name = field_name.replace("NonUS", "Non-US ").replace("Perm", "Permanent ").replace("Mail", "Mailing ").replace("TIN", "Tax ID ").replace("PC", "Postal Code").replace("TreatyClaim", "Treaty Claim ").replace("Percentage", "Rate").replace("IncomeType", "Income Type").replace("ArticleReason", "Article/Reason").replace("9aCountry", "Country").replace("YesNo", "Yes/No")
                                form_fields[clean_name] = str(value).strip()
                except Exception:
                    pass  # Form extraction is best-effort

                # Extract text from each page
                for page_num, page in enumerate(reader.pages, 1):
                    text = page.extract_text() or ""
                    if text.strip():
                        pages.append({
                            "page": page_num,
                            "content": text.strip()
                        })
        except Exception as e:
            raise RuntimeError(f"Error reading PDF: {str(e)}")

        full_text = "\n\n".join([p["content"] for p in pages])

        # Append form field values as structured data for the AI
        if form_fields:
            form_section = "\n\n=== FILLED-IN FORM FIELD VALUES ===\n"
            for field, value in form_fields.items():
                form_section += f"{field}: {value}\n"
            form_section += "=== END FORM FIELD VALUES ==="
            full_text += form_section

            # Also store form fields in metadata
            metadata["form_fields"] = form_fields

        return {
            "text": full_text,
            "pages": pages,
            "metadata": metadata,
            "format": "pdf"
        }

    @staticmethod
    def _extract_docx(path: Path) -> dict:
        """Extract text from DOCX file."""
        try:
            doc = Document(str(path))
            
            # Extract metadata
            metadata = {
                "title": doc.core_properties.title or "",
                "author": doc.core_properties.author or "",
                "paragraphs": len(doc.paragraphs)
            }
            
            # Extract text from paragraphs
            paragraphs = []
            for para in doc.paragraphs:
                if para.text.strip():
                    paragraphs.append(para.text.strip())
            
            full_text = "\n\n".join(paragraphs)
            
            # Create pages-like structure (group paragraphs)
            pages = []
            current_page = []
            page_num = 1
            
            for i, para in enumerate(paragraphs):
                current_page.append(para)
                # Create new "page" every 20 paragraphs or at natural breaks
                if len(current_page) >= 20 or para.endswith(("\n", ".", "!", "?")):
                    pages.append({
                        "page": page_num,
                        "content": "\n".join(current_page)
                    })
                    current_page = []
                    page_num += 1
            
            # Add remaining paragraphs
            if current_page:
                pages.append({
                    "page": page_num,
                    "content": "\n".join(current_page)
                })
            
            return {
                "text": full_text,
                "pages": pages,
                "metadata": metadata,
                "format": "docx"
            }
        except Exception as e:
            raise RuntimeError(f"Error reading DOCX: {str(e)}")

    @staticmethod
    def _extract_txt(path: Path) -> dict:
        """Extract text from plain text file."""
        try:
            with open(path, "r", encoding="utf-8") as file:
                text = file.read()
            
            # Split into pages-like chunks (by line count)
            lines = text.split("\n")
            pages = []
            page_num = 1
            current_page = []
            
            for line in lines:
                current_page.append(line)
                if len(current_page) >= 50:  # ~50 lines per page
                    pages.append({
                        "page": page_num,
                        "content": "\n".join(current_page)
                    })
                    current_page = []
                    page_num += 1
            
            if current_page:
                pages.append({
                    "page": page_num,
                    "content": "\n".join(current_page)
                })
            
            return {
                "text": text,
                "pages": pages,
                "metadata": {"lines": len(lines)},
                "format": "txt"
            }
        except UnicodeDecodeError:
            # Try with different encoding
            with open(path, "r", encoding="latin-1") as file:
                text = file.read()
            return {
                "text": text,
                "pages": [{"page": 1, "content": text}],
                "metadata": {"lines": len(text.split("\n"))},
                "format": "txt"
            }
